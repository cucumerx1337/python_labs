import csv
from pathlib import Path
import sys
import openpyxl

# Добавляем путь к корню репозитория
ROOT_DIR = Path(__file__).resolve().parent.parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.lib.tables import check_rows


def csv_to_xlsx(csv_path: str | Path, xlsx_path: str | Path) -> None:
    """Конвертирует CSV-файл в документ формата XLSX.

    Параметры:
        csv_path: путь к исходному CSV-файлу
        xlsx_path: путь для сохранения XLSX-файла

    Ошибки:
        FileNotFoundError: если исходный файл не существует
        ValueError: если файл пуст, содержит менее двух строк или содержит пустые ячейки
    """
    csv_target = Path(csv_path)
    xlsx_target = Path(xlsx_path)

    if not csv_target.exists():
        raise FileNotFoundError(f"Файл не найден: {csv_target}")
    if not csv_target.is_file():
        raise ValueError(f"Указанный путь не является файлом: {csv_target}")

    with csv_target.open("r", encoding="utf-8") as f:
        sample = f.read(2048)
        if not sample.strip():
            raise ValueError("CSV-файл не должен быть пустым")
        f.seek(0)

        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(f, dialect=dialect)
        rows = list(reader)

    if len(rows) <= 1:
        raise ValueError(
            "Недостаточно данных в CSV: требуется заголовок и как минимум одна строка с данными"
        )
    check_rows(rows)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 8)

    if not xlsx_target.parent.exists():
        xlsx_target.parent.mkdir(parents=True, exist_ok=True)

    wb.save(xlsx_target)


if __name__ == "__main__":
    in_csv = ROOT_DIR / "data/samples/people.csv"
    out_xlsx = ROOT_DIR / "data/out/people_from_csv.xlsx"

    print("=" * 60)

    try:
        print("Test 3: CSV -> XLSX")

        with open(in_csv, "r", encoding="utf-8") as f:
            initial_csv_len = len(f.readlines()) - 1

        csv_to_xlsx(in_csv, out_xlsx)

        wb = openpyxl.load_workbook(out_xlsx)
        ws = wb["Sheet1"]
        final_xlsx_len = ws.max_row - 1

        print(f" -> Zapisi CSV: {initial_csv_len}")
        print(f" -> Zapisi XLSX: {final_xlsx_len}")

        assert initial_csv_len == final_xlsx_len, "Error"
        print(" Success")
        print("=" * 60)

    except Exception as e:
        print(f"\nError: {e}")
