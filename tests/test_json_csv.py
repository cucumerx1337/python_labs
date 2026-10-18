import csv
import json
import sys
from pathlib import Path
import openpyxl

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

try:
    import pytest
except ImportError:
    pytest = None

from src.lab05.csv_xlsx import csv_to_xlsx
from src.lab05.json_csv import csv_to_json, json_to_csv


def test_json_to_csv_roundtrip(tmp_path):
    src = tmp_path / "people.json"
    dst = tmp_path / "people.csv"
    data = [
        {"name": "Alice", "age": 22},
        {"name": "Bob", "age": 25},
    ]

    src.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    json_to_csv(str(src), str(dst))

    with dst.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    assert len(rows) == 2
    assert {"name", "age"} <= set(rows[0].keys())


def test_csv_to_json_roundtrip(tmp_path):
    src = tmp_path / "people.csv"
    dst = tmp_path / "people.json"
    data = "name,age\nalice,22\nbob,25"

    src.write_text(data, encoding="utf-8")
    csv_to_json(str(src), str(dst))

    with dst.open(encoding="utf-8") as f:
        rows = json.load(f)

    assert len(rows) == 2
    assert {"name", "age"} <= set(rows[0].keys())


def test_csv_to_xlsx_roundtrip(tmp_path):
    src = tmp_path / "people.csv"
    dst = tmp_path / "people.xlsx"
    data = "name,age\nalice,22\nbob,25"

    src.write_text(data, encoding="utf-8")

    csv_to_xlsx(str(src), str(dst))

    wb = openpyxl.load_workbook(dst)
    ws = wb["Sheet1"]
    final_xlsx_len = ws.max_row - 1

    assert final_xlsx_len == 2
    assert {"name", "age"} <= set([cell.value for cell in ws[1]])


if __name__ == "__main__":
    import tempfile

    with tempfile.TemporaryDirectory() as td:
        t_path = Path(td)
        test_json_to_csv_roundtrip(t_path)
        test_csv_to_json_roundtrip(t_path)
        test_csv_to_xlsx_roundtrip(t_path)
    print("All tests passed successfully!")

